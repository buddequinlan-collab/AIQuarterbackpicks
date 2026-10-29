"""Build site/data.js from the project's slate workbooks and bet tracker.

Free picks are written in full. Premium picks are redacted to game + market
type, so the real pick never ships in the public site source; members get the
full card on Whop.

Usage: python3 export_site_data.py [--projections DIR] [--results DIR] [--out FILE]
"""
import argparse, json, os, re, datetime
from pathlib import Path
import openpyxl

ROOT = Path(os.environ.get("PROJECT_FILES", "/mnt/project-files"))  # slate workbooks + results live here
FREE_PER_SLATE = 1        # how many picks per slate are shown free
MAX_SLATES = 6            # newest slates to publish

SPORT_NAMES = {"nba": "NBA", "nfl": "NFL", "cfb": "College Football", "nhl": "NHL",
               "mlb": "MLB", "mnf": "NFL", "snf": "NFL", "tnf": "NFL"}

# Header aliases across the different workbook layouts
COLS = {
    "kind": ["Kind", "Type", "Call", "#"],
    "game": ["Game"],
    "pick": ["Pick", "Bet"],
    "odds": ["Price", "FD odds", "FanDuel odds", "Odds (FanDuel)", "Odds"],
    "model": ["Model %"],
    "book": ["No-vig %", "Book no-vig %", "Book %"],
    "edge": ["Edge"],
    "why": ["Why", "Why / risk"],
    "time": ["Kickoff (ET)", "Start"],
}
SKIP_KINDS = {"pass", "early"}


def pct(v):
    if v is None or v == "":
        return None
    v = float(v)
    return round(v * 100 if v <= 1 else v, 1)


def fmt_odds(v):
    if v is None:
        return None
    if isinstance(v, str):
        return v
    v = int(round(v))
    return f"+{v}" if v > 0 else str(v)


def market_of(pick):
    p = pick.lower()
    for key, label in [("anytime", "Anytime TD/goal"), ("first td", "First scorer"),
                       ("rec yds", "Receiving yards"), ("receiving", "Receiving yards"),
                       ("receptions", "Receptions"), ("rush", "Rushing yards"),
                       ("pass td", "Passing TDs"), ("pass yds", "Passing yards"),
                       ("pts+reb", "Points + rebounds"), ("pts+ast", "Points + assists"),
                       ("rebounds", "Rebounds"), ("assists", "Assists"), ("threes", "Threes"),
                       ("points", "Points"), (" k", "Strikeouts"), ("hits", "Hits"),
                       ("rbi", "RBIs"), ("shots", "Shots on goal"), ("spread", "Spread"),
                       ("ml", "Moneyline"), ("moneyline", "Moneyline"), ("team total", "Team total"),
                       ("over", "Total"), ("under", "Total")]:
        if key in p:
            return label
    if re.search(r"[+-]\d", p):
        return "Spread"
    return "Pick"


def read_named(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = next((n for n in wb.sheetnames if n.lower() == "named picks"), None)
    if not sheet:
        return None
    rows = list(wb[sheet].iter_rows(values_only=True))
    hdr = [str(h).strip() if h else "" for h in rows[0]]
    idx = {k: next((hdr.index(a) for a in al if a in hdr), None) for k, al in COLS.items()}
    get = lambda r, k: r[idx[k]] if idx[k] is not None and idx[k] < len(r) else None
    picks = []
    for r in rows[1:]:
        pick = get(r, "pick")
        if not pick:
            continue
        kind = str(get(r, "kind") or "Pick")
        if kind.lower() in SKIP_KINDS:
            continue
        tier = "lean" if "lean" in kind.lower() else "pick"
        game = get(r, "game")
        if not game and "," in str(pick):  # CFB layout: "Pick text, Away @ Home 3:30p"
            pick, game = [s.strip() for s in str(pick).split(",", 1)]
        raw_model = get(r, "model")
        model, book = pct(raw_model), pct(get(r, "book"))
        if book is None and model is not None and get(r, "edge") is not None:
            # NBA layout has model % and edge but no book %; edge shares model's scale
            scale = 100 if float(raw_model) <= 1 else 1
            book = round(model - float(get(r, "edge")) * scale, 1)
        picks.append({
            "tier": tier, "game": game, "pick": str(pick), "odds": fmt_odds(get(r, "odds")),
            "model_pct": model, "book_pct": book,
            "edge": round(model - book, 1) if model is not None and book is not None else None,
            "why": get(r, "why"), "time": get(r, "time"), "market": market_of(str(pick)),
        })
    return picks


def slate_from_file(path):
    m = re.match(r"([a-z]+)_.*?(\d{4}-\d{2}-\d{2})", path.name)
    if not m:
        return None
    picks = read_named(path)
    if not picks:
        return None
    # straight picks before leans, then by edge
    picks.sort(key=lambda p: (p["tier"] != "pick", -(p["edge"] or 0)))
    out = []
    for i, p in enumerate(picks):
        if i < FREE_PER_SLATE:
            out.append({**p, "locked": False})
        else:
            out.append({"locked": True, "tier": p["tier"], "game": p["game"],
                        "market": p["market"], "time": p["time"]})
    return {"id": path.stem, "sport": SPORT_NAMES.get(m.group(1), m.group(1).upper()),
            "date": m.group(2), "picks": out}


def record_from_tracker(path):
    ws = openpyxl.load_workbook(path, data_only=True)["Bets"]
    rows = list(ws.iter_rows(values_only=True))
    hdr = list(rows[0])
    c = {h: hdr.index(h) for h in hdr if h}
    bets = []
    for r in rows[1:]:
        res = r[c["Result"]]
        stake = r[c["Stake"]]
        if res not in ("Won", "Lost") or not stake:
            continue
        units = round(r[c["Profit"]] / stake, 2)  # each ticket normalized to 1u risked
        bets.append({"date": r[c["Date Placed"]], "sport": r[c["Sport"]], "type": r[c["Bet Type"]],
                     "game": r[c["Game(s)"]], "selections": r[c["Selections"]],
                     "odds": r[c["Final Odds"]], "result": "W" if res == "Won" else "L",
                     "units": units})
    bets.sort(key=lambda b: b["date"], reverse=True)
    w = sum(b["result"] == "W" for b in bets)
    return {"source": "Tracked tickets, 1 unit risked per ticket", "wins": w,
            "losses": len(bets) - w, "units": round(sum(b["units"] for b in bets), 2),
            "bets": bets}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--projections", default=ROOT / "projections", type=Path)
    ap.add_argument("--results", default=ROOT / "results", type=Path)
    ap.add_argument("--out", default=Path(__file__).resolve().parents[1] / "site" / "data.js", type=Path)
    a = ap.parse_args()
    slates = [s for f in a.projections.glob("*.xlsx") if (s := slate_from_file(f))]
    slates.sort(key=lambda s: s["date"], reverse=True)
    data = {"generated": datetime.datetime.now(datetime.timezone.utc).isoformat(timespec="minutes"),
            "slates": slates[:MAX_SLATES],
            "record": record_from_tracker(a.results / "bet_tracker.xlsx")}
    a.out.write_text("window.SITE_DATA = " + json.dumps(data, indent=1, default=str) + ";\n")
    print(f"wrote {a.out}: {len(data['slates'])} slates, record "
          f"{data['record']['wins']}-{data['record']['losses']} {data['record']['units']:+}u")


if __name__ == "__main__":
    main()
